"""Construit le jeu de fine-tuning de la synthèse par distillation depuis Claude.

Étapes (chaque étape est mise en cache dans `--out`, le script est reprenable) :

1. **Corpus disjoint** : Claude rédige des documents fictifs (contrats, rapports
   financiers, audits) avec des sociétés, montants et dates *différents* de ceux
   du corpus d'évaluation ; ils sont écrits en DOCX puis passent par la vraie
   chaîne ingestion → chunking → ChromaDB + BM25 (`HybridStore`).
2. **Questions** : Claude rédige, par document, des questions répondables et des
   questions plausibles mais sans réponse dans le document.
3. **Contexte réel** : chaque question est posée au retrieval hybride (k=5), sur
   tout le corpus d'entraînement : les exemples contiennent donc des extraits
   non pertinents et des échecs de retrieval, comme à l'inférence.
4. **Réponse cible** : Claude (enseignant) produit le `GroundedAnswer` avec le
   prompt système et le format d'extraits exacts du pipeline.

Sortie : `train.jsonl`, une ligne par exemple : `{"input": ..., "target": ...}`,
où `input` est exactement le texte que `HFLocalLLMClient` envoie au modèle.

    uv run python finetune/build_dataset.py --out finetune/data
"""

import argparse
import json
import re
import tempfile
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from pydantic import BaseModel, Field

from document_intelligence.chunking.service import chunk_document
from document_intelligence.core.config import settings
from document_intelligence.evaluation.runner import load_dataset
from document_intelligence.ingestion.service import ingest
from document_intelligence.orchestration.anthropic_client import AnthropicLLMClient
from document_intelligence.orchestration.cost import CostLedger
from document_intelligence.orchestration.llm import json_instruction_prompt
from document_intelligence.orchestration.models import GroundedAnswer
from document_intelligence.orchestration.prompts import (
    SYNTHESIS_SYSTEM_PROMPT,
    build_synthesis_prompt,
)
from document_intelligence.vectorization.chroma_store import ChromaStore
from document_intelligence.vectorization.hybrid_store import HybridStore

GOLDEN = Path("src/document_intelligence/evaluation/golden")

# Entités du corpus d'évaluation (fixtures + corpus étendu) : interdites.
FORBIDDEN = [
    "Helio Analytics",
    "HelioCore",
    "Verdier",
    "ACME",
    "Nordis",
    "Nordvale",
    "Tecnomar",
]

# Types de documents, calqués sur les cas d'usage du projet mais avec d'autres faits.
DOC_BRIEFS = [
    "contrat de prestation de services informatiques (infogérance) entre deux sociétés",
    "contrat de bail commercial pour des bureaux",
    "contrat-cadre de fourniture de composants industriels",
    "accord de confidentialité et de partenariat de recherche",
    "rapport financier annuel d'une PME industrielle (exercice 2024)",
    "rapport financier trimestriel d'une entreprise de logistique",
    "rapport financier semestriel d'une enseigne de distribution",
    "rapport d'audit qualité ISO 9001 d'un site de production",
    "rapport d'audit de conformité RGPD d'un éditeur de logiciels",
    "rapport d'audit de sécurité d'un hôpital privé",
]


class DocSection(BaseModel):
    heading: str
    paragraphs: list[str]


class DocSpec(BaseModel):
    title: str
    sections: list[DocSection] = Field(description="10 à 12 sections")
    table: list[list[str]] = Field(
        description="Tableau : ligne d'en-tête puis 3 à 5 lignes"
    )


class GeneratedQuestion(BaseModel):
    question: str
    answerable: bool = Field(
        description="False si le document ne permet pas de répondre"
    )
    evidence: str = Field(
        default="",
        description="Passage exact du document qui répond (vide si answerable=false)",
    )


class QuestionSet(BaseModel):
    questions: list[GeneratedQuestion]


def _doc_prompt(brief: str, index: int) -> str:
    return (
        f"Rédige un document d'entreprise fictif en français : {brief}.\n"
        "Contraintes :\n"
        "- 10 à 12 sections titrées (articles, chapitres...), 1 à 3 paragraphes chacune ;\n"
        "- beaucoup de faits précis et vérifiables : montants, pourcentages, dates, "
        "durées, noms de sociétés et de personnes, obligations, seuils ;\n"
        "- un tableau chiffré (ligne d'en-tête + 3 à 5 lignes) ;\n"
        "- sociétés, personnes, produits et chiffres entièrement inventés ; "
        f"n'utilise aucun de ces noms : {', '.join(FORBIDDEN)} ;\n"
        f"- document n°{index + 1} d'une série : choisis des noms originaux."
    )


def _questions_prompt(doc: DocSpec) -> str:
    body = "\n".join(
        f"## {s.heading}\n" + "\n".join(s.paragraphs) for s in doc.sections
    )
    table = "\n".join(" | ".join(row) for row in doc.table)
    return (
        f"<document>\n# {doc.title}\n{body}\n\nTableau :\n{table}\n</document>\n\n"
        "Rédige 30 questions qu'un analyste poserait sur ce document :\n"
        "- 24 questions répondables, variées (chiffre, date, durée, obligation, "
        "liste, comparaison entre deux sections, tableau), formulées comme un "
        "utilisateur réel (sans recopier les phrases du document) ; `evidence` = "
        "passage exact qui répond ;\n"
        "- 6 questions plausibles pour ce type de document mais dont la réponse "
        "n'y figure PAS (answerable=false, evidence vide)."
    )


def _write_docx(doc: DocSpec, path: Path) -> None:
    document = Document()
    document.add_heading(doc.title, level=0)
    for section in doc.sections:
        document.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            document.add_paragraph(paragraph)
    rows = [r for r in doc.table if r]
    if rows:
        width = max(len(r) for r in rows)
        table = document.add_table(rows=len(rows), cols=width)
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                table.cell(i, j).text = value
    document.save(str(path))


def _cached(path: Path, schema: type[BaseModel], produce):  # type: ignore[no-untyped-def]
    if path.exists():
        return schema.model_validate_json(path.read_text(encoding="utf-8"))
    value = produce()
    path.write_text(value.model_dump_json(indent=1), encoding="utf-8")
    return value


def leakage_report(docs: list[DocSpec], questions: list[str]) -> dict[str, object]:
    """Recouvrements entre le corpus d'entraînement et le jeu d'évaluation."""
    text = "\n".join(
        " ".join([d.title, *(p for s in d.sections for p in s.paragraphs)])
        for d in docs
    ).lower()
    eval_examples = [
        ex
        for name in ("fixtures_qa.jsonl", "extended_qa.jsonl")
        for ex in load_dataset(GOLDEN / name)
    ]
    substrings = {s for ex in eval_examples for s in ex.relevant_substrings}
    forbidden_hits = [name for name in FORBIDDEN if name.lower() in text]
    substring_hits = [s for s in substrings if len(s) >= 12 and s.lower() in text]
    near_duplicates = []
    for ex in eval_examples:
        best = max(
            questions,
            key=lambda q: SequenceMatcher(None, q.lower(), ex.question.lower()).ratio(),
        )
        ratio = SequenceMatcher(None, best.lower(), ex.question.lower()).ratio()
        if ratio >= 0.8:
            near_duplicates.append((ex.question, best, round(ratio, 2)))
    return {
        "forbidden_entities_found": forbidden_hits,
        "eval_relevant_substrings_found": substring_hits,
        "near_duplicate_questions_ratio_ge_0.8": near_duplicates,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=Path("finetune/data"))
    parser.add_argument("--k", type=int, default=5)
    args = parser.parse_args()
    out: Path = args.out
    (out / "docs").mkdir(parents=True, exist_ok=True)
    (out / "questions").mkdir(exist_ok=True)

    key = settings.anthropic_api_key
    if not key:
        raise SystemExit("ANTHROPIC_API_KEY requis (enseignant Claude).")
    ledger = CostLedger()
    writer = AnthropicLLMClient(
        key, settings.anthropic_model, max_tokens=8000, role="dataset", ledger=ledger
    )
    teacher = AnthropicLLMClient(
        key, settings.anthropic_model, role="teacher", ledger=ledger
    )

    docs: list[DocSpec] = []
    for i, brief in enumerate(DOC_BRIEFS):
        spec = _cached(
            out / "docs" / f"doc{i:02d}.json",
            DocSpec,
            lambda b=brief, i=i: writer.complete_structured(_doc_prompt(b, i), DocSpec),
        )
        docs.append(spec)
        print(f"document {i}: {spec.title} ({len(spec.sections)} sections)")

    question_sets = [
        _cached(
            out / "questions" / f"doc{i:02d}.json",
            QuestionSet,
            lambda d=doc: writer.complete_structured(_questions_prompt(d), QuestionSet),
        )
        for i, doc in enumerate(docs)
    ]

    all_questions = [q for qs in question_sets for q in qs.questions]
    report = leakage_report(docs, [q.question for q in all_questions])
    (out / "leakage_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    print("contrôle de fuite :", json.dumps(report, ensure_ascii=False))

    train_path = out / "train.jsonl"
    done = set()
    if train_path.exists():
        done = {
            json.loads(line)["question"]
            for line in train_path.read_text(encoding="utf-8").splitlines()
        }

    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        chroma = ChromaStore(persist_directory=str(Path(tmp) / "chroma"))
        for i, doc in enumerate(docs):
            path = Path(tmp) / f"train_doc{i:02d}.docx"
            _write_docx(doc, path)
            chroma.add_chunks(chunk_document(ingest(path)))
        store = HybridStore(chroma)

        with train_path.open("a", encoding="utf-8") as sink:
            for q in all_questions:
                if q.question in done:
                    continue
                retrieved = store.query(q.question, k=args.k)
                prompt = build_synthesis_prompt(q.question, retrieved)
                target = teacher.complete_structured(
                    prompt, GroundedAnswer, system=SYNTHESIS_SYSTEM_PROMPT
                )
                # Comme le nœud de synthèse : seules les citations présentes dans
                # le contexte sont conservées dans la cible.
                known = {r.chunk_id for r in retrieved}
                target = target.model_copy(
                    update={
                        "cited_chunk_ids": [
                            c for c in target.cited_chunk_ids if c in known
                        ]
                    }
                )
                evidence_retrieved = bool(q.evidence) and any(
                    re.sub(r"\s+", " ", q.evidence.lower())[:60]
                    in re.sub(r"\s+", " ", r.text.lower())
                    for r in retrieved
                )
                sink.write(
                    json.dumps(
                        {
                            "question": q.question,
                            "answerable": q.answerable,
                            "evidence_retrieved": evidence_retrieved,
                            "input": json_instruction_prompt(
                                prompt, GroundedAnswer, SYNTHESIS_SYSTEM_PROMPT
                            ),
                            "target": target.model_dump_json(),
                        },
                        ensure_ascii=False,
                    )
                    + "\n"
                )
                sink.flush()

    rows = [
        json.loads(line) for line in train_path.read_text(encoding="utf-8").splitlines()
    ]
    targets = [GroundedAnswer.model_validate_json(r["target"]) for r in rows]
    print(f"exemples : {len(rows)}")
    print(f"  répondables : {sum(r['answerable'] for r in rows)}")
    print(
        f"  passage attendu dans le contexte : {sum(r['evidence_retrieved'] for r in rows)}"
    )
    print(
        f"  cibles avec abstention : {sum(not t.sufficient_context for t in targets)}"
    )
    print(
        "coût enseignant :", ledger.summary(settings.anthropic_model)["cost_actual_usd"]
    )


if __name__ == "__main__":
    main()
